from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
TEXT = (ROOT / 'opening.txt').read_text(encoding='utf-8') + (ROOT / 'activities-and-ending.txt').read_text(encoding='utf-8')

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = 'Malgun Gothic'
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    s.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
doc.styles['Title'].font.size = Pt(24)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)
for line in TEXT.splitlines():
    if not line.strip():
        continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        p = doc.add_paragraph(line[3:], 'Heading 1')
        p.paragraph_format.space_before = Pt(16)
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], 'Heading 2')
    else:
        p = doc.add_paragraph()
        if ': “' in line and line.split(':', 1)[0] in ['모아', '소망이', '차곡', '봉봉', '토리']:
            who, speech = line.split(':', 1)
            p.add_run(who + '  ').bold = True
            p.add_run(speech.strip())
        else:
            r = p.add_run(line)
            if line.startswith('['):
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor.from_string('555555')
        p.paragraph_format.widow_control = True
        p.paragraph_format.keep_together = True
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('포포리 마을 초반 대사집  ·  ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.save(ROOT / '포포리_마을_초반_대사집_v1.docx')
print(ROOT / '포포리_마을_초반_대사집_v1.docx')
